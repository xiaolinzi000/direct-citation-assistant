# -*- coding: utf-8 -*-
"""边界场景回归：中文 key / 小写 [cite:] / author-year 无年份 / 标题后正文段保护 / .doc 校验。
用法：python edge_test.py（自建临时素材，测完清理）"""
import io
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
SCRIPTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
INSERT = os.path.join(SCRIPTS, "insert_refs.py")

TMP = tempfile.mkdtemp(prefix="edge_")
DOCX = os.path.join(TMP, "文稿.docx")
DOCX2 = os.path.join(TMP, "保护.docx")
REFS = os.path.join(TMP, "refs.csv")

REFS_DATA = [
    {"key": "注意力机制", "type": "journal", "title": "注意力机制研究综述",
     "authors": "赵六, 孙七", "source": "计算机学报", "year": "2024",
     "volume": "47", "issue": "1", "pages": "1-20", "doi": "", "url": "", "note": ""},
    {"key": "smith2020", "type": "journal", "title": "Deep learning survey",
     "authors": "Smith, J., Doe, A.", "source": "Nature", "year": "2020",
     "volume": "580", "issue": "", "pages": "1-10", "doi": "", "url": "", "note": ""},
    {"key": "noyear", "type": "journal", "title": "A classic work",
     "authors": "Zhang, W.", "source": "Proc. AAAI", "year": "",
     "volume": "", "issue": "", "pages": "", "doi": "", "url": "", "note": ""},
]


def w(tag):
    return f"{{{W}}}{tag}"


def run_insert(args):
    r = subprocess.run([sys.executable, INSERT, *args],
                       capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    return r.stdout + r.stderr


def para_text(p):
    return "".join(t.text or "" for t in p.iter(w("t")))


def build_docx(path):
    from docx import Document
    doc = Document()
    doc.add_heading("边界测试文稿", level=1)
    p1 = doc.add_paragraph()
    p1.add_run("中文文献支撑该观点[CITE:注意力机制]。")
    p2 = doc.add_paragraph()
    p2.add_run("小写占位符也可识别[cite:smith2020]。")
    p3 = doc.add_paragraph()
    p3.add_run("无年份引用[CITE:noyear]。")
    doc.save(path)


def build_protected_docx(path):
    """构造：参考文献标题之后紧跟一段以 [9] 开头的正文（无书签）——不应被误删。"""
    from docx import Document
    doc = Document()
    doc.add_heading("保护测试", level=1)
    doc.add_paragraph("正文内容[?]。")
    doc.add_heading("参考文献", level=2)
    doc.add_paragraph("[9] 这是参考文献标题之后的正文段落，不应被脚本删除。")
    doc.save(path)


def write_refs():
    with io.open(REFS, "w", encoding="utf-8-sig", newline="") as f:
        import csv
        wtr = csv.DictWriter(f, fieldnames=list(REFS_DATA[0].keys()))
        wtr.writeheader()
        for r in REFS_DATA:
            wtr.writerow(r)


def read_body(path):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(w("body"))
    hyps = [h.get(w("anchor")) for h in body.iter(w("hyperlink"))
            if (h.get(w("anchor")) or "").startswith("ref_")]
    return body, hyps


def main():
    write_refs()
    build_docx(DOCX)
    ok = True

    # 测试1：编号制 + 中文 key + 小写 cite + 无年份
    out = run_insert(["--docx", DOCX, "--refs", REFS])
    body, hyps = read_body(DOCX)
    exp = ["ref_注意力机制", "ref_smith2020", "ref_noyear"]
    assert hyps == exp, f"超链接异常: {hyps}（期望 {exp}）"
    print(f"✅ 测试1 中文 key / 小写 [cite:] / 占位符 → 3 个超链接 {hyps}")

    # 测试2：author-year 无年份输出 (Zhang)（不带逗号空年）
    out = run_insert(["--docx", DOCX, "--refs", REFS, "--citation", "author-year"])
    body, hyps = read_body(DOCX)
    labels = []
    with zipfile.ZipFile(DOCX) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    for h in root.iter(w("hyperlink")):
        txt = "".join(t.text or "" for t in h.iter(w("t"))).strip()
        if txt:
            labels.append(txt)
    assert "(赵六等, 2024)" in labels, f"中文 author-year 异常: {labels}"
    assert "(Smith et al., 2020)" in labels, f"英文 author-year 异常: {labels}"
    assert "(Zhang)" in labels, f"无年份应为 (Zhang): {labels}"
    print(f"✅ 测试2 author-year 标签: {labels}")

    # 测试3：参考文献标题后的正文段保护
    build_protected_docx(DOCX2)
    out = run_insert(["--docx", DOCX2, "--refs", REFS])
    assert "请先 add_refs" not in out  # [?] 取 refs 第一条，正常
    with zipfile.ZipFile(DOCX2) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(w("body"))
    texts = [para_text(p) for p in body.iter(w("p"))]
    assert any("[9] 这是参考文献标题之后的正文段落" in t for t in texts), \
        f"正文段被误删！段落: {texts}"
    print("✅ 测试3 参考文献标题后的正文段未被误删")

    # 测试4：.doc 扩展名校验
    fake_doc = os.path.join(TMP, "旧格式.doc")
    with open(fake_doc, "w") as f:
        f.write("x")
    out = run_insert(["--docx", fake_doc, "--refs", REFS])
    assert "只支持 .docx" in out, f".doc 应被拒绝: {out}"
    print("✅ 测试4 .doc 扩展名校验生效")

    shutil.rmtree(TMP, ignore_errors=True)
    print("\n🎉 全部边界测试通过")


if __name__ == "__main__":
    main()
