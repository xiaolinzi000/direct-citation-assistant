# -*- coding: utf-8 -*-
"""author-year 同年同作者后缀测试（GB/T 7714 著者-出版年制）：
同（作者, 年份）的多篇文献按题名字母序加 a/b 后缀，正文与文末一致，重跑幂等。
用法：python year_suffix_test.py（自建素材，不依赖 make_demo）"""
import csv
import io
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree
from docx import Document

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
INSERT = os.path.join(SCRIPTS, "insert_refs.py")

tmp = tempfile.mkdtemp(prefix="ys_")
REFS = os.path.join(tmp, "refs.csv")
DOCX = os.path.join(tmp, "文.docx")


def w(tag):
    return f"{{{W}}}{tag}"


def para_text(p):
    return "".join(t.text or "" for t in p.iter(w("t")))


def make_refs():
    fields = ["key", "type", "title", "authors", "source", "year",
              "volume", "issue", "pages", "city", "doi", "url", "note"]
    rows = [
        {"key": "zhang_b", "type": "journal", "title": "Beta 方法研究",
         "authors": "张伟, 李静", "source": "计算机学报", "year": "2023",
         "note": "支持句1：方法对比"},
        {"key": "zhang_a", "type": "journal", "title": "Alpha 机制综述",
         "authors": "张伟, 李静", "source": "软件学报", "year": "2023",
         "note": "支持句2：机制综述"},
        {"key": "li2022", "type": "journal", "title": "模型压缩方法",
         "authors": "李静", "source": "电子学报", "year": "2022",
         "note": "支持句3：压缩方法"},
    ]
    with io.open(REFS, "w", encoding="utf-8-sig", newline="") as f:
        wr = csv.DictWriter(f, fieldnames=fields)
        wr.writeheader()
        for r in rows:
            wr.writerow(r)


def make_docx():
    doc = Document()
    doc.add_paragraph("张伟等在 2023 年提出了 Beta 方法[CITE:zhang_b]，"
                      "其机制在 Alpha 综述中有系统总结[CITE:zhang_a]。")
    doc.add_paragraph("模型压缩的早期工作见[CITE:li2022]。")
    doc.save(DOCX)


def run_insert():
    r = subprocess.run([sys.executable, INSERT, "--docx", DOCX,
                        "--refs", REFS, "--citation", "author-year"],
                       capture_output=True, text=True, encoding="utf-8",
                       errors="replace")
    if r.returncode != 0:
        print(r.stdout + r.stderr)
        sys.exit(1)
    return r.stdout


def summary():
    with zipfile.ZipFile(DOCX) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(w("body"))
    hyps = [(h.get(w("anchor")) or "", para_text(h).strip())
            for h in body.iter(w("hyperlink"))]
    entries = []
    for p in body.iter(w("p")):
        for bs in p.iter(w("bookmarkStart")):
            nm = (bs.get(w("name")) or "")
            if nm.startswith("ref_") and nm not in entries:
                entries.append(nm)
    # 文末条目文本（含年份）
    entry_txt = {}
    for p in body.iter(w("p")):
        keys = [(bs.get(w("name")) or "") for bs in p.iter(w("bookmarkStart"))
                if (bs.get(w("name")) or "").startswith("ref_")]
        if keys:
            entry_txt[keys[0]] = para_text(p).strip()
    return hyps, entry_txt


try:
    make_refs()
    make_docx()
    print("== 首次运行 ==")
    print(run_insert().strip().splitlines()[-2:])
    hyps, entries = summary()
    body_labels = [t for _, t in hyps if t.startswith("(")]
    print("正文引用:", body_labels)
    print("文末条目:")
    for k, v in entries.items():
        print(f"  {k}: {v[:70]}")

    # 断言1：正文同年同作者带 a/b 后缀
    assert "(张伟等, 2023a)" in body_labels, f"缺 2023a: {body_labels}"
    assert "(张伟等, 2023b)" in body_labels, f"缺 2023b: {body_labels}"
    # 断言2：文末条目年份带后缀（题名字母序 Beta 后置）→ 2023a=Alpha, 2023b=Beta
    ta = next(v for k, v in entries.items() if k == "ref_zhang_a")
    tb = next(v for k, v in entries.items() if k == "ref_zhang_b")
    assert "2023a" in ta, f"Alpha 应为 2023a: {ta}"
    assert "2023b" in tb, f"Beta 应为 2023b: {tb}"
    # 断言3：不同作者不同年不加后缀
    tl = next(v for k, v in entries.items() if k == "ref_li2022")
    assert "2022" in tl and "2022a" not in tl, f"li2022 不应有后缀: {tl}"
    print("✅ 断言1-3 通过（正文 a/b / 文末年份后缀 / 无后缀不误加）")

    # 重跑幂等
    print("\n== 重跑（幂等）==")
    print(run_insert().strip().splitlines()[-2:])
    hyps2, entries2 = summary()
    assert [t for _, t in hyps2 if t.startswith("(")] == body_labels, "重跑后正文标签变化！"
    assert entries2 == entries, "重跑后文末条目变化！"
    print("✅ 幂等通过")
    print("🎉 year_suffix 测试全部通过")
finally:
    shutil.rmtree(tmp, ignore_errors=True)
