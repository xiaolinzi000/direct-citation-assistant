# -*- coding: utf-8 -*-
"""分章/多文档测试：跨文档统一编号 + 参考文献表在主文档 + 跨文档跳转链接。
用法：先 make_demo.py 生成 refs.csv，再跑本脚本（自建 第1章/第2章 docx）。"""
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from lxml import etree

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
P_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
SCRIPTS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
HERE = os.path.dirname(os.path.abspath(__file__))
INSERT = os.path.join(SCRIPTS, "insert_refs.py")

TMP = tempfile.mkdtemp(prefix="multi_")
CH1 = os.path.join(TMP, "第1章.docx")
CH2 = os.path.join(TMP, "第2章.docx")
REFS = os.path.join(HERE, "refs.csv")


def w(t):
    return f"{{{W}}}{t}"


def para_text(p):
    return "".join(t.text or "" for t in p.iter(w("t")))


def build_chapter(path, paras):
    from docx import Document
    doc = Document()
    doc.add_heading(os.path.basename(path).replace(".docx", ""), level=1)
    for text in paras:
        doc.add_paragraph(text)
    doc.save(path)


def run_insert():
    r = subprocess.run([sys.executable, INSERT, "--docx", CH1, CH2,
                        "--refs", REFS],
                       capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
    return r.stdout + r.stderr


def inspect(path):
    with zipfile.ZipFile(path) as z:
        root = etree.fromstring(z.read("word/document.xml"))
        rels_xml = z.read("word/_rels/document.xml.rels") \
            if "word/_rels/document.xml.rels" in z.namelist() else None
    rels_root = etree.fromstring(rels_xml) if rels_xml else None
    hyps = []
    for h in root.iter(w("hyperlink")):
        anchor = h.get(w("anchor")) or ""
        if anchor.startswith("ref_"):
            txt = "".join(t.text or "" for t in h.iter(w("t"))).strip()
            rid = h.get(f"{{{R_NS}}}id")
            hyps.append((anchor, txt, rid))
    external = []
    if rels_root is not None:
        for rel in rels_root:
            if rel.get("TargetMode") == "External" and \
               "hyperlink" in (rel.get("Type") or ""):
                external.append((rel.get("Id"), rel.get("Target")))
    return hyps, external


def main():
    build_chapter(CH1, ["注意力机制综述[?]。", "预训练模型相关工作[CITE:devlin2019]。"])
    build_chapter(CH2, ["大模型最新进展[CITE:zhang2023]。"])
    out = run_insert()
    print(out.splitlines()[-8:])

    # 断言 1：第一章编号 1,2；第二章编号 3
    h1, ext1 = inspect(CH1)
    h2, ext2 = inspect(CH2)
    t1 = [t for _, t, _ in h1]
    t2 = [t for _, t, _ in h2]
    assert t1 == ["[1]", "[2]"], f"第一章编号异常: {t1}"
    assert t2 == ["[3]"], f"第二章编号异常: {t2}"
    print(f"✅ 跨文档统一编号：第1章 {t1} / 第2章 {t2}")

    # 断言 2：第一章超链接带 r:id 且 rels 指向第2章（跨文档）；第二章不带 r:id
    rids1 = [r for _, _, r in h1]
    rids2 = [r for _, _, r in h2]
    assert all(rids1), f"第一章超链接缺跨文档 r:id: {rids1}"
    assert not any(rids2), f"第二章（主文档）不应有跨文档 r:id: {rids2}"
    assert ext1, f"第一章 rels 缺 External 链接: {ext1}"
    targets = [t for _, t in ext1]
    assert all("第2章.docx" in t for t in targets), f"跨文档目标错误: {targets}"
    print(f"✅ 跨文档跳转：第1章 {len(ext1)} 个链接 → {targets[0]}")

    # 断言 3：文末表在主文档（第二章）3 条
    with zipfile.ZipFile(CH2) as z:
        root = etree.fromstring(z.read("word/document.xml"))
    entries = [para_text(p)[:40] for p in root.iter(w("p"))
               if any((bs.get(w("name")) or "").startswith("ref_")
                      for bs in p.iter(w("bookmarkStart")))]
    assert len(entries) == 3, f"主文档条目数异常: {entries}"
    assert entries[0].startswith("[1]") and entries[2].startswith("[3]")
    print(f"✅ 主文档（第2章）文末表 {len(entries)} 条：{[e[:20] for e in entries]}")

    # 断言 4：重跑幂等（编号不变、跨文档链接不重复）
    out2 = run_insert()
    assert "编号连续性验证通过" in out2
    h1b, ext1b = inspect(CH1)
    assert [t for _, t, _ in h1b] == t1, "重跑后第一章编号变化！"
    assert len(ext1b) == len(ext1), f"重跑后跨文档链接重复: {len(ext1b)} vs {len(ext1)}"
    print("✅ 重跑幂等：编号不变、跨文档链接不重复")

    shutil.rmtree(TMP, ignore_errors=True)
    print("\n🎉 分章/多文档测试全部通过")


if __name__ == "__main__":
    main()
